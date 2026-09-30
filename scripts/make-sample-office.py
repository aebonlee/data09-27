# 예시 첨부(워드·파워포인트·엑셀) 만들기 — 내용은 모두 가상입니다.
#   pip install python-docx python-pptx openpyxl
#   python3 scripts/make-sample-office.py
# 결과: samples/collect/att/M0001/캡_1차시안_검토회의록.docx
#       samples/collect/att/M0003/CMF_샘플_평가표.xlsx
#       samples/collect/att/M0005/전시부스_렌더링_v2.pptx
# 그다음 node scripts/make-samples.js 가 이 파일들의 글을 뽑아 manifest.json 을 만듭니다.
import os
import docx
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PPt
import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'samples', 'collect', 'att')


def out(mid, name):
    d = os.path.join(ROOT, mid)
    os.makedirs(d, exist_ok=True)
    return os.path.join(d, name)


# ── 워드: 회의록 (문단 + 표 + 줄바꿈 · 탭 · 특수 글자) ──
d = docx.Document()
d.styles['Normal'].font.size = Pt(10.5)
d.add_heading('캡 인테리어 개선 — 1차 시안 검토 회의록 (가상)', level=1)
t = d.add_table(rows=3, cols=2)
t.style = 'Table Grid'
for i, (k, v) in enumerate([('일시', '2026-09-21 09:00'), ('참석', '디자인팀장(가상), 보고자(가상), 설계 담당(가상)'), ('안건', '1차 시안 3종 비교 & 최종안 선정')]):
    t.cell(i, 0).text = k
    t.cell(i, 1).text = v
d.add_paragraph('결정 사항')
d.add_paragraph('B안 최종 시안 선정 완료.')
p = d.add_paragraph('조작부 버튼 배열은 협력사 치수 회신 후 확정할 예정입니다.')
d.add_paragraph('이슈')
d.add_paragraph('협력사 조작부 치수 회신이 지연되고 있어 확인이 필요합니다.')
p = d.add_paragraph('다음 할 일')
p.add_run().add_break()
p.add_run('3D 모델링 업데이트\t10/2까지')
d.save(out('M0001', '캡_1차시안_검토회의록.docx'))

# ── 엑셀: 평가표 2개 시트 (공유 문자열 · 숫자 · 빈 칸) ──
wb = openpyxl.Workbook()
ws = wb.active
ws.title = '평가표'
ws.append(['번호', '샘플', '색차(ΔE)', '판정', '비고'])
rows = [(1, '딥블루 무광', 0.8, '합격', ''), (2, '딥블루 유광', 1.1, '합격', ''), (3, '웜그레이', 2.9, '재도장 필요', '기준 2.0 초과'),
        (4, '쿨그레이', 1.4, '합격', ''), (5, '옐로 포인트', 3.2, '재도장 필요', '기준 2.0 초과')]
for r in rows:
    ws.append(list(r))
ws2 = wb.create_sheet('일정')
ws2.append(['항목', '일자'])
ws2.append(['재도장 샘플 재입고 예정', '2026-09-30'])
ws2.append(['평가 완료 목표', '2026-10-02'])
wb.save(out('M0003', 'CMF_샘플_평가표.xlsx'))

# ── 파워포인트: 3장 (제목 · 글상자 · 줄바꿈) ──
pr = Presentation()
s = pr.slides.add_slide(pr.slide_layouts[0])
s.shapes.title.text = '전시 부스 렌더링 v2 (가상)'
s.placeholders[1].text = '배포일: 2026-09-24'
s = pr.slides.add_slide(pr.slide_layouts[1])
s.shapes.title.text = '주요 변경'
tf = s.placeholders[1].text_frame
tf.text = '조명 위치 조정 완료'
tf.add_paragraph().text = '관람 동선 안 반영 완료'
s = pr.slides.add_slide(pr.slide_layouts[5])
s.shapes.title.text = '남은 일'
tb = s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame
tb.text = '부스 그래픽 시안 확정 예정(10/5)'
tb.paragraphs[0].runs[0].font.size = PPt(20)
pr.save(out('M0005', '전시부스_렌더링_v2.pptx'))
print('ok: docx · xlsx · pptx')
